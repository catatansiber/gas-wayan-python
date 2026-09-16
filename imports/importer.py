"""Orkestrasi importer F3: verifikasi checksum, baca sheet Database read-only lewat
openpyxl, tulis RawRow immutable + StagingRow. Tidak pernah menulis catalog/operations
(master/lifecycle) - lihat docs/phase-plan.md F3 dan CLAUDE.md batas perubahan."""

import hashlib
from pathlib import Path

import openpyxl
from django.db import transaction
from django.utils import timezone

from imports import parsing
from imports.exceptions import ChecksumMismatch, SheetNotFound
from imports.models import ImportBatch, RawRow, StagingRow

SHEET_NAME = "Database"

# Checksum SHA-256 SERI TABUNG.xlsm terverifikasi F0 - lihat CLAUDE.md dan docs/source-audit.md.
EXPECTED_WORKBOOK_SHA256 = "D96FC4E7366BFC9B38CBD2BEB6985CAB361BCA9E1D36CFAED5DD1FFDDFA33B81"


class _DryRunRollback(Exception):
    pass


def compute_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def run_import(
    workbook_path,
    expected_sha256: str,
    actor=None,
    dry_run: bool = False,
    sheet_name: str = SHEET_NAME,
) -> ImportBatch:
    workbook_path = Path(workbook_path)
    actual_sha256 = compute_sha256(workbook_path)
    if actual_sha256.lower() != expected_sha256.lower():
        raise ChecksumMismatch(expected_sha256, actual_sha256)

    batch = ImportBatch(
        workbook_filename=workbook_path.name,
        workbook_sha256=actual_sha256,
        sheet_name=sheet_name,
        parser_version=parsing.PARSER_VERSION,
        dry_run=dry_run,
        triggered_by=actor,
    )

    try:
        with transaction.atomic():
            batch.save()
            wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                if sheet_name not in wb.sheetnames:
                    raise SheetNotFound(sheet_name)
                ws = wb[sheet_name]
                _import_rows(ws, batch)
            finally:
                wb.close()
            batch.status = ImportBatch.Status.COMPLETED
            batch.finished_at = timezone.now()
            batch.save()
            if dry_run:
                raise _DryRunRollback
    except _DryRunRollback:
        pass
    except Exception as exc:
        batch.status = ImportBatch.Status.FAILED
        batch.error_message = str(exc)
        raise
    return batch


def _import_rows(ws, batch: ImportBatch) -> None:
    for row_number, raw_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(raw_values[:6])
        values += [None] * (6 - len(values))

        batch.raw_rows_seen += 1
        is_blank = all(parsing.trim(v) == "" for v in values)
        if is_blank:
            batch.blank_rows_seen += 1

        raw_row, raw_created = _get_or_create_raw_row(batch, row_number, values, is_blank)
        if raw_created:
            batch.raw_rows_created += 1
        else:
            batch.raw_rows_already_existing += 1

        if is_blank:
            continue

        _, staging_created = _get_or_create_staging_row(batch, raw_row, values)
        if staging_created:
            batch.staging_rows_created += 1
        else:
            batch.staging_rows_already_existing += 1


_COLUMN_KEYS = (
    "nomor",
    "nomor_tabung",
    "nama_relasi",
    "tanggal_pengiriman",
    "tanggal_pengembalian",
    "jenis_gas",
)


def _get_or_create_raw_row(batch, row_number, values, is_blank):
    existing = RawRow.objects.filter(
        workbook_sha256=batch.workbook_sha256,
        sheet_name=batch.sheet_name,
        source_row_number=row_number,
    ).first()
    if existing is not None:
        return existing, False

    raw_json = {
        key: {"value": parsing.cell_to_json(value), "type": parsing.cell_type_label(value)}
        for key, value in zip(_COLUMN_KEYS, values, strict=True)
    }
    raw_row = RawRow.objects.create(
        workbook_sha256=batch.workbook_sha256,
        sheet_name=batch.sheet_name,
        source_row_number=row_number,
        is_blank=is_blank,
        raw_json=raw_json,
        first_seen_batch=batch,
    )
    return raw_row, True


def _get_or_create_staging_row(batch, raw_row, values):
    existing = StagingRow.objects.filter(
        raw_row=raw_row, parser_version=parsing.PARSER_VERSION
    ).first()
    if existing is not None:
        return existing, False

    legacy_nomor_raw = parsing.trim(values[0])
    serial_raw = parsing.trim(values[1])
    customer_raw = parsing.trim(values[2])
    gas_raw = parsing.trim(values[5])

    sent = parsing.parse_date_cell(values[3])
    returned = parsing.classify_return_cell(values[4])
    gas_code, gas_known = parsing.match_gas_code(gas_raw) if gas_raw else ("", False)

    warnings: list[str] = []
    errors: list[str] = []

    if not serial_raw:
        errors.append("SERIAL_KOSONG")

    if not sent.raw_text:
        warnings.append("TANGGAL_KIRIM_KOSONG")
    elif not sent.parseable:
        warnings.append("TANGGAL_KIRIM_TIDAK_TERPARSE")

    ReturnStatus = StagingRow.ReturnStatus
    if returned.status == ReturnStatus.NONE:
        warnings.append("TANGGAL_KEMBALI_KOSONG")
    elif returned.status == ReturnStatus.OTHER_TEXT:
        warnings.append("STATUS_KEMBALI_TIDAK_DIKENAL")
    elif returned.status == ReturnStatus.SDH_BALIK:
        warnings.append("STATUS_SDH_BALIK_BELUM_DIPETAKAN")
    elif returned.date_unknown:
        warnings.append("TANGGAL_KEMBALI_TIDAK_DIKETAHUI")

    if not gas_raw:
        warnings.append("JENIS_GAS_KOSONG")
    elif not gas_known:
        warnings.append("JENIS_GAS_TIDAK_DIKENAL")

    if errors:
        review_status = StagingRow.ReviewStatus.ERROR
    elif warnings:
        review_status = StagingRow.ReviewStatus.NEEDS_REVIEW
    else:
        review_status = StagingRow.ReviewStatus.READY

    dedupe_key = (
        parsing.build_dedupe_key(serial_raw, sent.raw_text, returned.raw_text) if serial_raw else ""
    )

    staging = StagingRow.objects.create(
        raw_row=raw_row,
        parser_version=parsing.PARSER_VERSION,
        legacy_nomor_raw=legacy_nomor_raw,
        serial_number_raw=serial_raw,
        serial_number_present=bool(serial_raw),
        customer_name_raw=customer_raw,
        customer_name_normalized=parsing.normalize_customer_name(customer_raw)
        if customer_raw
        else "",
        sent_at_raw=sent.raw_text,
        sent_at=sent.value,
        sent_at_is_native=sent.is_native,
        sent_at_parseable=sent.parseable,
        returned_at_raw=returned.raw_text,
        returned_at=returned.date,
        returned_at_is_native=returned.is_native,
        returned_status=returned.status,
        returned_date_unknown=returned.date_unknown,
        gas_code_raw=gas_raw,
        gas_type_code=gas_code,
        gas_type_known=gas_known,
        dedupe_key=dedupe_key,
        warnings=warnings,
        errors=errors,
        review_status=review_status,
        import_batch=batch,
    )
    return staging, True

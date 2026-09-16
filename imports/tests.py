import datetime
import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
from django.test import TestCase

from catalog.models import Customer, Cylinder, GasType
from identity.models import Role, User
from imports import importer, parsing, publisher
from imports.exceptions import ChecksumMismatch
from imports.models import ImportBatch, PublishBatch, RawRow, ReviewDecision, SourceLink, StagingRow
from operations import services
from operations.exceptions import DomainError
from operations.models import Cycle, LifecycleEvent


def _build_workbook(rows, sheet_name="Database", extra_sheet="SearchData"):
    """rows: list baris, masing-masing list 6 nilai (A..F) atau None untuk baris kosong
    (tidak ditulis sama sekali - mensimulasikan baris tanpa elemen <row> di XML asli)."""
    tmp_dir = Path(tempfile.mkdtemp())
    path = tmp_dir / "synthetic.xlsm"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["Nomor", "Nomor Tabung", "Nama Relasi", "Tgl Kirim", "Tgl Kembali", "Jenis Gas"])
    for i, row in enumerate(rows, start=2):
        if row is None:
            continue
        for col, value in enumerate(row, start=1):
            ws.cell(row=i, column=col, value=value)

    other = wb.create_sheet(extra_sheet)
    other.append(["JANGAN_DIBACA"])
    other.sheet_state = "hidden"

    wb.save(path)
    return path


class ParsingUnitTests(TestCase):
    def test_native_date_cell(self):
        result = parsing.parse_date_cell(datetime.date(2025, 9, 8))
        self.assertTrue(result.is_native)
        self.assertEqual(result.value, datetime.date(2025, 9, 8))

    def test_text_date_day_first(self):
        result = parsing.parse_date_cell("08/09/2025")
        self.assertFalse(result.is_native)
        self.assertEqual(result.value, datetime.date(2025, 9, 8))

    def test_text_date_two_digit_year(self):
        result = parsing.parse_date_cell("8.9.25")
        self.assertEqual(result.value, datetime.date(2025, 9, 8))

    def test_impossible_date_not_guessed(self):
        result = parsing.parse_date_cell("29/2/25")
        self.assertIsNone(result.value)
        self.assertFalse(result.parseable)

    def test_unrecognized_format_not_guessed(self):
        result = parsing.parse_date_cell("September 8 2025")
        self.assertIsNone(result.value)

    def test_returned_keyword_k(self):
        result = parsing.classify_return_cell("k")
        self.assertEqual(result.status, StagingRow.ReturnStatus.K)
        self.assertTrue(result.date_unknown)
        self.assertIsNone(result.date)

    def test_returned_keyword_kembali(self):
        result = parsing.classify_return_cell("KEMBALI")
        self.assertEqual(result.status, StagingRow.ReturnStatus.KEMBALI)

    def test_returned_keyword_sdh_kembali(self):
        result = parsing.classify_return_cell("Sdh Kembali")
        self.assertEqual(result.status, StagingRow.ReturnStatus.SDH_KEMBALI)

    def test_sdh_balik_not_mapped_to_returned(self):
        result = parsing.classify_return_cell("SDH BALIK")
        self.assertEqual(result.status, StagingRow.ReturnStatus.SDH_BALIK)
        self.assertFalse(result.date_unknown)

    def test_other_free_text_needs_review(self):
        result = parsing.classify_return_cell("GANTI TABUNG")
        self.assertEqual(result.status, StagingRow.ReturnStatus.OTHER_TEXT)

    def test_gas_code_official(self):
        code, known = parsing.match_gas_code("o2")
        self.assertEqual(code, "O2")
        self.assertTrue(known)

    def test_gas_code_unofficial_not_guessed(self):
        code, known = parsing.match_gas_code("BBS")
        self.assertEqual(code, "")
        self.assertFalse(known)

    def test_csv_formula_injection_sanitized(self):
        self.assertEqual(parsing.sanitize_csv_cell("=cmd()"), "'=cmd()")
        self.assertEqual(parsing.sanitize_csv_cell("PT Aman"), "PT Aman")


class ImportWorkbookTests(TestCase):
    def _run(self, rows, expected_sha256=None):
        path = _build_workbook(rows)
        sha256 = expected_sha256 or importer.compute_sha256(path)
        return path, importer.run_import(path, expected_sha256=sha256)

    def test_checksum_mismatch_writes_nothing(self):
        path = _build_workbook([["1", "TB-0001", "Budi", "08/09/2025", "K", "O2"]])
        with self.assertRaises(ChecksumMismatch):
            importer.run_import(path, expected_sha256="0" * 64)
        self.assertEqual(RawRow.objects.count(), 0)
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_only_database_sheet_is_read(self):
        _, batch = self._run(
            [["1", "TB-0001", "Budi", "08/09/2025", "K", "O2"]],
        )
        self.assertEqual(batch.raw_rows_seen, 1)
        self.assertFalse(RawRow.objects.filter(raw_json__nomor__value="JANGAN_DIBACA").exists())

    def test_blank_row_counted_explicitly(self):
        rows = [
            ["1", "TB-0001", "Budi", "08/09/2025", "K", "O2"],
            None,
            ["3", "TB-0003", "Citra", "10/09/2025", "", ""],
        ]
        _, batch = self._run(rows)
        self.assertEqual(batch.raw_rows_seen, 3)
        self.assertEqual(batch.blank_rows_seen, 1)
        self.assertEqual(RawRow.objects.filter(is_blank=True).count(), 1)
        self.assertEqual(StagingRow.objects.count(), 2)

    def test_leading_zero_serial_preserved(self):
        self._run([["1", "007", "Budi", "08/09/2025", "K", "O2"]])
        staging = StagingRow.objects.get()
        self.assertEqual(staging.serial_number_raw, "007")

    def test_mixed_date_formats(self):
        rows = [
            ["1", "TB-0001", "Budi", datetime.date(2025, 9, 8), "K", "O2"],
            ["2", "TB-0002", "Citra", "9.9.25", "K", "O2"],
            ["3", "TB-0003", "Dedi", "31/2/25", "K", "O2"],
        ]
        self._run(rows)
        by_serial = {s.serial_number_raw: s for s in StagingRow.objects.all()}
        self.assertTrue(by_serial["TB-0001"].sent_at_is_native)
        self.assertEqual(by_serial["TB-0002"].sent_at, datetime.date(2025, 9, 9))
        self.assertFalse(by_serial["TB-0003"].sent_at_parseable)
        self.assertEqual(by_serial["TB-0003"].review_status, StagingRow.ReviewStatus.NEEDS_REVIEW)

    def test_missing_serial_is_error_status(self):
        self._run([["1", "", "Budi", "08/09/2025", "K", "O2"]])
        staging = StagingRow.objects.get()
        self.assertEqual(staging.review_status, StagingRow.ReviewStatus.ERROR)
        self.assertIn("SERIAL_KOSONG", staging.errors)

    def test_unofficial_gas_flagged_needs_review_not_guessed(self):
        self._run([["1", "TB-0001", "Budi", "08/09/2025", "K", "BBS"]])
        staging = StagingRow.objects.get()
        self.assertFalse(staging.gas_type_known)
        self.assertEqual(staging.gas_type_code, "")
        self.assertEqual(staging.gas_code_raw, "BBS")
        self.assertEqual(staging.review_status, StagingRow.ReviewStatus.NEEDS_REVIEW)

    def test_rerun_same_snapshot_is_idempotent(self):
        rows = [
            ["1", "TB-0001", "Budi", "08/09/2025", "K", "O2"],
            ["2", "TB-0002", "Citra", "09/09/2025", "", ""],
        ]
        path = _build_workbook(rows)
        sha256 = importer.compute_sha256(path)

        batch1 = importer.run_import(path, expected_sha256=sha256)
        self.assertEqual(batch1.raw_rows_created, 2)
        self.assertEqual(RawRow.objects.count(), 2)
        self.assertEqual(StagingRow.objects.count(), 2)

        batch2 = importer.run_import(path, expected_sha256=sha256)
        self.assertEqual(batch2.raw_rows_created, 0)
        self.assertEqual(batch2.raw_rows_already_existing, 2)
        self.assertEqual(batch2.staging_rows_already_existing, 2)
        self.assertEqual(RawRow.objects.count(), 2)
        self.assertEqual(StagingRow.objects.count(), 2)

    def test_dry_run_writes_nothing(self):
        path = _build_workbook([["1", "TB-0001", "Budi", "08/09/2025", "K", "O2"]])
        sha256 = importer.compute_sha256(path)
        batch = importer.run_import(path, expected_sha256=sha256, dry_run=True)
        self.assertEqual(batch.raw_rows_created, 1)
        self.assertEqual(RawRow.objects.count(), 0)
        self.assertEqual(StagingRow.objects.count(), 0)
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_last_rows_import_with_correct_source_row_numbers(self):
        rows = [["1", "TB-0001", "Budi", "08/09/2025", "K", "O2"]] * 3
        path = _build_workbook(rows)
        sha256 = importer.compute_sha256(path)
        importer.run_import(path, expected_sha256=sha256)
        numbers = sorted(RawRow.objects.values_list("source_row_number", flat=True))
        self.assertEqual(numbers, [2, 3, 4])

    def test_duplicate_candidate_key_ignores_legacy_nomor(self):
        rows = [
            ["100", "TB-0009", "Budi", "08/09/2025", "K", "O2"],
            ["999", "TB-0009", "Budi", "08/09/2025", "K", "O2"],
        ]
        self._run(rows)
        keys = set(StagingRow.objects.values_list("dedupe_key", flat=True))
        self.assertEqual(len(keys), 1)

    def test_alias_candidate_same_normalized_different_raw(self):
        rows = [
            ["1", "TB-0001", "Budi  Santoso", "08/09/2025", "K", "O2"],
            ["2", "TB-0002", "budi santoso", "09/09/2025", "K", "O2"],
        ]
        self._run(rows)
        normalized = set(StagingRow.objects.values_list("customer_name_normalized", flat=True))
        raw = set(StagingRow.objects.values_list("customer_name_raw", flat=True))
        self.assertEqual(len(normalized), 1)
        self.assertEqual(len(raw), 2)


class PublisherTestCaseMixin:
    """Data sintetis - tidak ada data pelanggan/tabung asli. Satu StagingRow di sini SELALU
    dibuat langsung lewat ORM (bukan lewat importer.run_import) supaya tes fokus pada logika
    publish per kelompok, bukan parsing F3 (sudah dites terpisah di atas)."""

    @classmethod
    def make_common_fixtures(cls):
        cls.admin = User.objects.create_user(username="f4admin", password="x", role=Role.ADMIN)
        cls.operator = User.objects.create_user(
            username="f4operator", password="x", role=Role.OPERATOR
        )
        cls.viewer = User.objects.create_user(username="f4viewer", password="x", role=Role.VIEWER)
        GasType.objects.get_or_create(code=GasType.Code.O2)
        cls.import_batch = ImportBatch.objects.create(
            workbook_filename="synthetic.xlsm",
            workbook_sha256="f" * 64,
            sheet_name="Database",
            parser_version=parsing.PARSER_VERSION,
        )

    def _make_staging(
        self,
        serial,
        source_row_number,
        sent_at=None,
        returned_at=None,
        review_status=StagingRow.ReviewStatus.READY,
        gas_code="O2",
        customer_raw="Budi Santoso",
    ):
        raw_row = RawRow.objects.create(
            workbook_sha256=self.import_batch.workbook_sha256,
            sheet_name="Database",
            source_row_number=source_row_number,
            is_blank=False,
            raw_json={},
            first_seen_batch=self.import_batch,
        )
        returned_status = (
            StagingRow.ReturnStatus.DATE if returned_at else StagingRow.ReturnStatus.NONE
        )
        return StagingRow.objects.create(
            raw_row=raw_row,
            parser_version=parsing.PARSER_VERSION,
            serial_number_raw=serial,
            serial_number_present=True,
            customer_name_raw=customer_raw,
            customer_name_normalized=parsing.normalize_customer_name(customer_raw),
            sent_at_raw=str(sent_at) if sent_at else "",
            sent_at=sent_at,
            sent_at_is_native=True,
            sent_at_parseable=sent_at is not None,
            returned_at_raw=str(returned_at) if returned_at else "",
            returned_at=returned_at,
            returned_at_is_native=bool(returned_at),
            returned_status=returned_status,
            gas_code_raw=gas_code,
            gas_type_code=gas_code if gas_code else "",
            gas_type_known=bool(gas_code),
            review_status=review_status,
            import_batch=self.import_batch,
        )


class PublishGroupTests(PublisherTestCaseMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.make_common_fixtures()

    def test_clean_single_row_group_publishes(self):
        row = self._make_staging(
            "TB-1001", 100, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        batch = publisher.run_publish(self.admin)

        row.refresh_from_db()
        self.assertEqual(row.publish_outcome, StagingRow.PublishOutcome.PUBLISHED)
        self.assertEqual(batch.groups_published, 1)
        self.assertEqual(batch.rows_published, 1)

        link = SourceLink.objects.get(staging_row=row)
        self.assertEqual(link.cycle.status, Cycle.Status.CLOSED)
        self.assertEqual(link.cycle.close_reason, Cycle.CloseReason.RETURNED)
        self.assertEqual(Cylinder.objects.get(serial_number="TB-1001").status, "AVAILABLE")
        self.assertTrue(
            LifecycleEvent.objects.filter(
                cylinder__serial_number="TB-1001", source=LifecycleEvent.Source.IMPORT
            ).exists()
        )

    def test_group_blocked_when_any_row_not_ready(self):
        row_ok = self._make_staging(
            "TB-1002", 101, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        row_bad = self._make_staging(
            "TB-1002",
            102,
            datetime.date(2025, 1, 6),
            None,
            review_status=StagingRow.ReviewStatus.NEEDS_REVIEW,
        )
        batch = publisher.run_publish(self.admin)

        row_ok.refresh_from_db()
        row_bad.refresh_from_db()
        self.assertEqual(row_ok.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)
        self.assertEqual(row_bad.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)
        self.assertEqual(batch.groups_blocked, 1)
        self.assertEqual(SourceLink.objects.count(), 0)

    def test_business_duplicate_blocks_until_excluded(self):
        row1 = self._make_staging(
            "TB-1003", 103, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        row2 = self._make_staging(
            "TB-1003", 104, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        batch = publisher.run_publish(self.admin)
        self.assertEqual(batch.groups_blocked, 1)
        row1.refresh_from_db()
        row2.refresh_from_db()
        self.assertEqual(row1.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)
        self.assertEqual(row2.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)

        publisher.exclude_staging_row(self.admin, row2, reason="Duplikat entri legacy")
        row2.refresh_from_db()
        self.assertEqual(row2.publish_outcome, StagingRow.PublishOutcome.EXCLUDED_APPROVED)
        self.assertEqual(ReviewDecision.objects.filter(staging_row=row2).count(), 1)

        batch2 = publisher.run_publish(self.admin)
        self.assertEqual(batch2.groups_published, 1)
        row1.refresh_from_db()
        self.assertEqual(row1.publish_outcome, StagingRow.PublishOutcome.PUBLISHED)

    def test_rerun_after_publish_is_idempotent(self):
        self._make_staging("TB-1004", 105, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5))
        batch1 = publisher.run_publish(self.admin)
        self.assertEqual(batch1.groups_published, 1)

        batch2 = publisher.run_publish(self.admin)
        self.assertEqual(batch2.groups_published, 0)
        self.assertEqual(batch2.groups_already_resolved, 1)
        self.assertEqual(SourceLink.objects.count(), 1)
        self.assertEqual(Cycle.objects.count(), 1)

    def test_crash_mid_batch_then_resume(self):
        self._make_staging("TB-1005", 106, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5))
        row_second_group = self._make_staging(
            "TB-1006", 107, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        row_third_group = self._make_staging(
            "TB-1007", 108, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )

        with self.assertRaises(RuntimeError):
            publisher.run_publish(self.admin, _fail_after_groups=1)

        row_second_group.refresh_from_db()
        row_third_group.refresh_from_db()
        published_now = StagingRow.objects.filter(
            publish_outcome=StagingRow.PublishOutcome.PUBLISHED
        ).count()
        self.assertEqual(published_now, 1, "kelompok pertama harus tetap commit walau crash")
        self.assertEqual(row_second_group.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)
        failed_batch = PublishBatch.objects.filter(status=PublishBatch.Status.FAILED).get()
        self.assertEqual(failed_batch.groups_published, 1)

        resumed = publisher.run_publish(self.admin)
        self.assertEqual(resumed.groups_published, 2)
        self.assertEqual(
            StagingRow.objects.filter(publish_outcome=StagingRow.PublishOutcome.PUBLISHED).count(),
            3,
        )
        self.assertEqual(SourceLink.objects.count(), 3)

    def test_group_rolls_back_completely_on_mid_group_failure(self):
        clean_row = self._make_staging(
            "TB-2000", 200, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        flaky_row1 = self._make_staging(
            "TB-3000", 201, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        flaky_row2 = self._make_staging(
            "TB-3000", 202, datetime.date(2025, 1, 6), datetime.date(2025, 1, 10)
        )

        real_return = services.return_cylinder

        def flaky(*args, **kwargs):
            if kwargs.get("idempotency_key") == f"f4-return-{flaky_row2.id}":
                raise RuntimeError("simulasi gagal di tengah kelompok (hanya untuk tes)")
            return real_return(*args, **kwargs)

        with patch("imports.publisher.services.return_cylinder", side_effect=flaky):
            with self.assertRaises(RuntimeError):
                publisher.run_publish(self.admin)

        clean_row.refresh_from_db()
        flaky_row1.refresh_from_db()
        flaky_row2.refresh_from_db()
        self.assertEqual(clean_row.publish_outcome, StagingRow.PublishOutcome.PUBLISHED)
        self.assertEqual(flaky_row1.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)
        self.assertEqual(flaky_row2.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)
        self.assertEqual(
            SourceLink.objects.filter(staging_row__in=[flaky_row1, flaky_row2]).count(), 0
        )
        self.assertEqual(Cycle.objects.filter(cylinder__serial_number="TB-3000").count(), 0)
        self.assertEqual(Cylinder.objects.filter(serial_number="TB-3000").exists(), False)

    def test_pending_gas_policy_blocks_group(self):
        row = self._make_staging(
            "TB-4000",
            300,
            datetime.date(2025, 1, 1),
            datetime.date(2025, 1, 5),
            review_status=StagingRow.ReviewStatus.NEEDS_REVIEW,
            gas_code="",
        )
        batch = publisher.run_publish(self.admin)
        row.refresh_from_db()
        self.assertEqual(batch.groups_blocked, 1)
        self.assertEqual(row.publish_outcome, StagingRow.PublishOutcome.PENDING_REVIEW)

    def test_role_enforcement_on_publish_and_exclude(self):
        self._make_staging("TB-5000", 400, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5))
        with self.assertRaises(DomainError):
            publisher.run_publish(self.viewer)

        row = self._make_staging(
            "TB-5001", 401, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        with self.assertRaises(DomainError):
            publisher.exclude_staging_row(self.operator, row, reason="Coba operator")
        with self.assertRaises(DomainError):
            publisher.exclude_staging_row(self.admin, row, reason="   ")

    def test_cannot_exclude_already_published_row(self):
        row = self._make_staging(
            "TB-6000", 500, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5)
        )
        publisher.run_publish(self.admin)
        row.refresh_from_db()
        with self.assertRaises(DomainError):
            publisher.exclude_staging_row(self.admin, row, reason="Terlambat")

    def test_reconciliation_raw_equals_published_plus_pending_plus_excluded(self):
        self._make_staging("TB-7000", 600, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5))
        self._make_staging(
            "TB-7001",
            601,
            datetime.date(2025, 1, 1),
            None,
            review_status=StagingRow.ReviewStatus.NEEDS_REVIEW,
        )
        row_to_exclude = self._make_staging(
            "TB-7002",
            602,
            datetime.date(2025, 1, 1),
            None,
            review_status=StagingRow.ReviewStatus.ERROR,
        )
        publisher.exclude_staging_row(self.admin, row_to_exclude, reason="Data rusak")
        publisher.run_publish(self.admin)

        total = StagingRow.objects.count()
        published = StagingRow.objects.filter(
            publish_outcome=StagingRow.PublishOutcome.PUBLISHED
        ).count()
        pending = StagingRow.objects.filter(
            publish_outcome=StagingRow.PublishOutcome.PENDING_REVIEW
        ).count()
        excluded = StagingRow.objects.filter(
            publish_outcome=StagingRow.PublishOutcome.EXCLUDED_APPROVED
        ).count()
        self.assertEqual(total, published + pending + excluded)
        self.assertEqual(published, 1)
        self.assertEqual(pending, 1)
        self.assertEqual(excluded, 1)

    def test_raw_archive_remains_searchable_after_publish(self):
        self._make_staging("TB-8000", 700, datetime.date(2025, 1, 1), datetime.date(2025, 1, 5))
        publisher.run_publish(self.admin)
        self.assertEqual(RawRow.objects.filter(source_row_number=700).count(), 1)
        staging = StagingRow.objects.get(raw_row__source_row_number=700)
        self.assertEqual(staging.serial_number_raw, "TB-8000")
        self.assertEqual(staging.publish_outcome, StagingRow.PublishOutcome.PUBLISHED)

    def test_customer_alias_exact_normalized_reused_across_rows(self):
        self._make_staging(
            "TB-9000",
            800,
            datetime.date(2025, 1, 1),
            datetime.date(2025, 1, 5),
            customer_raw="Budi  Santoso",
        )
        self._make_staging(
            "TB-9001",
            801,
            datetime.date(2025, 2, 1),
            datetime.date(2025, 2, 5),
            customer_raw="budi santoso",
        )
        publisher.run_publish(self.admin)
        self.assertEqual(Customer.objects.count(), 1)
        customer = Customer.objects.get()
        self.assertEqual(customer.aliases.count(), 1)
